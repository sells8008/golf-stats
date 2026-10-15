import openpyxl,csv,datetime,re
from openpyxl.styles import Font
wb=openpyxl.load_workbook('p5_src.xlsx')
NORM={'TPC Sawgrass - STADIUM COURSE':'TPC Sawgrass STADIUM COURSE','Talbot Country Club':'Talbot','Soldier Hollow - Gold Course':'Soldier Hollow'}
G=list(csv.DictReader(open('ghin_differentials.csv')))
for g in G: g['dt']=datetime.datetime.strptime(g['Date'],'%m/%d/%Y'); g['C']=NORM.get(g['Course'],g['Course']); g['rid']=None
rep=[]
for sn,rt in (('18-Hole Rounds','18'),('9-Hole Rounds','9')):
    ws=wb[sn];H=[c.value for c in ws[1]]
    if 'Differential' not in H: ws.cell(1,len(H)+1,'Differential').font=Font(bold=True);H.append('Differential')
    c={h:i+1 for i,h in enumerate(H)}
    for i in range(2,ws.max_row+1):
        dt=ws.cell(i,1).value
        if not dt: continue
        gross=ws.cell(i,c['Gross']).value;note=ws.cell(i,c['Notes']).value or '';course=ws.cell(i,2).value
        cand=[g for g in G if g['dt']==dt and g['Holes']==rt and not g['rid']]
        m=[g for g in cand if int(g['Score'])==gross]
        how='score'
        if not m:
            adj=re.search(r'adj (\d+)',note); m=[g for g in cand if adj and int(g['Score'])==int(adj.group(1))]; how='adj score (from Notes)'
        if not m:
            m=[g for g in cand if course.split()[0].lower() in g['Course'].lower()]; how='date+course only (GHIN adjusted score)'
            if len(m)!=1: m=[]
        if len(m)>1:
            m2=[g for g in m if course.split()[0].lower() in g['Course'].lower()]; m=m2 or m
        rid=ws.cell(i,c['RoundId']).value
        if len(m)==1:
            g=m[0];g['rid']=rid;ws.cell(i,c['Differential'],float(g['Differential']))
            if how!='score': rep.append(f'{rid}: matched on {how} {g["Score"]} (raw {gross}) -> {g["Differential"]}')
        else: rep.append(f'{rid}: NO MATCH ({len(cand)} same-date candidates: {[ (x["Score"],x["Course"]) for x in cand]})')
if 'GHIN Scores' in wb.sheetnames: del wb['GHIN Scores']
ws=wb.create_sheet('GHIN Scores');ws.append(['Date','Course','Tees','Holes','Score','Differential','RoundId'])
for x in ws[1]: x.font=Font(bold=True)
for g in sorted(G,key=lambda g:g['dt'],reverse=True):
    ws.append([g['dt'],g['C'],g['Tees'],int(g['Holes']),int(g['Score']),float(g['Differential']),g['rid']]);ws.cell(ws.max_row,1).number_format='yyyy-mm-dd'
ws.freeze_panes='A2'
wb.save('final/Hunter_Handicap_Tracker.xlsx')
print(len(G),'ghin rows;',sum(1 for g in G if g['rid']),'linked');print('\n'.join(rep))
