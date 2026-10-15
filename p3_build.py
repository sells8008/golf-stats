import openpyxl,datetime,collections
from openpyxl.styles import Font
wb=openpyxl.load_workbook('p3_src.xlsx')
H=wb['Holes']; hdr=[c.value for c in H[1]]; c={h:i+1 for i,h in enumerate(hdr)}
rows=list(range(2,H.max_row+1)); V=lambda r,k:H.cell(r,c[k]).value
by=collections.defaultdict(list)
for r in rows:
    if V(r,'RoundId'): by[V(r,'RoundId')].append(r)
full={rid for rid,rs in by.items() if any(V(r,'Putts') is not None for r in rs)}
nc=collections.Counter()
for rid in full:
    for r in by[rid]:
        if V(r,'Score') is None: continue
        if not V(r,'GIR'): H.cell(r,c['GIR'],'Miss'); H.cell(r,c['GreenMiss'],'No Chance'); nc['green']+=1
        if V(r,'Par')>=4 and not V(r,'FIR'): H.cell(r,c['FIR'],'Miss'); H.cell(r,c['FairwayMiss'],'No Chance'); nc['fw']+=1
        att=1 if V(r,'GIR')=='Miss' else 0
        H.cell(r,c['UpDownAttempt'],att); H.cell(r,c['UpDownMade'],1 if att and V(r,'Score')<=V(r,'Par') else 0)
print('No Chance set:',dict(nc))
# add missing rounds
R=wb['18-Hole Rounds']; rh=[x.value for x in R[1]]; rc={h:i+1 for i,h in enumerate(rh)}
have={R.cell(i,rc['RoundId']).value for i in range(2,R.max_row+1)}
notes={'2026-06-26_park-city-gc_18':'Rating/Slope/Yards blank: public Blue listings conflict (70.0/133/6215 BlueGolf+MyGolfIndex vs 69.7/129/6277 GolfPass vs 69.3/130 Nomads); hole pars match BlueGolf card. Verify in GHIN.',
       '2026-06-27_canyons-golf_18':'Rating/Slope/Yards blank: public Blue listings conflict (66.3/134/5522 GolfPass/TeeOff vs 65.8/130/5317 Stymie) and hole par sequence differs from published cards (routing change?). Verify in GHIN.'}
for rid,note in notes.items():
    if rid in have: continue
    rs=by[rid]; i=R.max_row+1
    while R.cell(i-1,1).value is None: i-=1
    R.insert_rows(i) if False else None
    dt=V(rs[0],'RoundDate'); par=sum(V(r,'Par') for r in rs); gross=sum(V(r,'Score') for r in rs)
    tp=[V(r,'Score')-V(r,'Par') for r in rs]
    vals={'Date':dt,'Course':V(rs[0],'Course'),'Tees':V(rs[0],'Tees'),'Par':par,'Gross':gross,'Doubles (exact +2)':sum(x==2 for x in tp),'TriplesOrWorse (+3 or more)':sum(x>=3 for x in tp),'Notes':'GHIN hole-by-hole (par & score only); '+note,'RoundId':rid}
    for k,v in vals.items(): R.cell(i,rc[k],v)
    R.cell(i,1).number_format='yyyy-mm-dd'
    R.cell(i,rc['ToPar'],f'=IF(COUNT(G{i}:H{i})=2,H{i}-G{i},"")')
    R.cell(i,rc['DoublesOrWorse (+2 or more)'],f'=IF(COUNT(J{i}:K{i})=2,J{i}+K{i},"")')
    R.cell(i,rc['BlowUpIndex'],f'=IF(COUNT(J{i}:K{i})=2,J{i}+2*K{i},"")')
    R.cell(i,rc['FIR_Pct'],f'=IF(COUNT(N{i}:O{i})=2,N{i}/O{i},"")')
    R.cell(i,rc['Last20 (auto)'],f'=IF(A{i}="","",IF(COUNTIFS($A$2:$A$1000,">"&A{i})+COUNTIFS(A{i}:$A$1000,A{i})-1<20,1,0))')
    print('added',rid,par,gross,vals['Doubles (exact +2)'],vals['TriplesOrWorse (+3 or more)'])
# per-round fills
for sn in ('18-Hole Rounds','9-Hole Rounds'):
    S=wb[sn]; sh=[x.value for x in S[1]]; sc={h:i+1 for i,h in enumerate(sh)}
    for i in range(2,S.max_row+1):
        rid=S.cell(i,sc['RoundId']).value
        if rid not in full: continue
        rs=by[rid]
        S.cell(i,sc['GIR'],sum(V(r,'GIR')=='Hit' for r in rs))
        S.cell(i,sc['UpDownsAttempted'],sum(V(r,'UpDownAttempt') or 0 for r in rs)); S.cell(i,sc['UpDownsMade'],sum(V(r,'UpDownMade') or 0 for r in rs))
        f=[r for r in rs if V(r,'Par')>=4]
        S.cell(i,sc['FairwaysHit'],sum(V(r,'FIR')=='Hit' for r in f)); S.cell(i,sc['FairwaysPossible'],len(f))
        print(rid,'GIR',S.cell(i,sc['GIR']).value,'UD',S.cell(i,sc['UpDownsMade']).value,'/',S.cell(i,sc['UpDownsAttempted']).value,'FW',S.cell(i,sc['FairwaysHit']).value)
# regenerate Fairway Holes
fh=wb['Fairway Holes']; fh.delete_rows(2,fh.max_row); i=2
L=sorted((r for r in rows if V(r,'RoundDate')),key=lambda r:(V(r,'RoundDate'),V(r,'Hole')))
for r in L:
    if V(r,'Par')>=4 and V(r,'FIR') in ('Hit','Miss') and V(r,'Score') is not None:
        fh.cell(i,1,V(r,'RoundDate')).number_format='yyyy-mm-dd'; fh.cell(i,2,V(r,'Course')); fh.cell(i,3,f"{V(r,'RoundType')}-hole"); fh.cell(i,4,V(r,'Hole')); fh.cell(i,5,V(r,'Par')); fh.cell(i,6,V(r,'Score'))
        fh.cell(i,7,f'=IF(COUNT(E{i}:F{i})=2,F{i}-E{i},"")'); fh.cell(i,8,V(r,'FIR'))
        fh.cell(i,9,f'=IF(C{i}="9-hole",1,IF(COUNTIFS(\'18-Hole Rounds\'!$A$2:$A$1000,A{i},\'18-Hole Rounds\'!$R$2:$R$1000,1)>0,1,0))'); i+=1
print('fairway rows',i-2)
import os; os.makedirs('final',exist_ok=True); wb.save('final/Hunter_Handicap_Tracker.xlsx')
