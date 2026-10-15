import openpyxl
from openpyxl.styles import Font
wb=openpyxl.load_workbook('p1_src.xlsx')
tees={}
for sn,rt in (('18-Hole Rounds',18),('9-Hole Rounds',9)):
    ws=wb[sn]; mc=ws.max_column
    for i,h in enumerate(['GIR','Putts','ThreePuttsPlus','OnePutts','UpDownsMade','UpDownsAttempted']):
        c=ws.cell(1,mc+1+i,h); c.font=Font(bold=True)
    for r in ws.iter_rows(min_row=2,values_only=True):
        if r[0]: tees[(r[0],rt)]=r[2]
if 'Holes' in wb.sheetnames: del wb['Holes']
ws=wb.create_sheet('Holes')
hdr=['RoundDate','Course','Tees','RoundType','Hole','Par','Score','ToPar','FIR','FairwayMiss','GIR','GreenMiss','Putts','UpDownAttempt','UpDownMade']
ws.append(hdr)
for c in ws[1]: c.font=Font(bold=True)
fh=wb['Fairway Holes']; n=0
rows=[r for r in fh.iter_rows(min_row=2,values_only=True) if r[0]]
rows.sort(key=lambda r:(r[0],r[3]))
for r in rows:
    rt=9 if r[2]=='9-hole' else 18
    n+=1; i=n+1
    ws.append([r[0],r[1],tees.get((r[0],rt)),rt,r[3],r[4],r[5],f'=IF(COUNT(F{i}:G{i})=2,G{i}-F{i},"")',r[7],None,None,None,None,None,None])
    ws.cell(i,1).number_format='yyyy-mm-dd'
ws.freeze_panes='A2'
wb.save('Hunter_Handicap_Tracker_p1.xlsx'); print(n,'rows')
