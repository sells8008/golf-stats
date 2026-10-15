import openpyxl,csv,datetime,collections,re
from openpyxl.styles import Font
wb=openpyxl.load_workbook('up/Hunter_Handicap_Tracker.xlsx')
CSV=list(csv.DictReader(open('ghin_holes.csv')))
oldFH=[r for r in wb['Fairway Holes'].iter_rows(min_row=2,values_only=True) if r[0]]
g=collections.OrderedDict()
for r in CSV:
    r['dt']=datetime.datetime.strptime(r['RoundDate'],'%m/%d/%Y'); g.setdefault((r['dt'],r['RoundType']),[]).append(r)
# 9/19: split two 9-hole rounds (18 rows) by score order -> gross 42/45
k919=(datetime.datetime(2026,9,19),'9'); rows=g.pop(k919)
a,b=rows[:9],rows[9:]
if sum(int(x['Score']) for x in a)!=42: a,b=b,a
assert sum(int(x['Score']) for x in a)==42 and sum(int(x['Score']) for x in b)==45, [x['Hole'] for x in rows]
g[(k919[0],'9','42')]=a; g[(k919[0],'9','45')]=b
slug=lambda s:re.sub(r'[^a-z0-9]+','-',s.lower()).strip('-')[:20]
# round sheet index
RS={}
for sn,rt in (('18-Hole Rounds','18'),('9-Hole Rounds','9')):
    ws=wb[sn]; H=[c.value for c in ws[1]]
    if 'RoundId' not in H: ws.cell(1,len(H)+1,'RoundId').font=Font(bold=True); H.append('RoundId')
    ci={h:i+1 for i,h in enumerate(H)}
    for row in range(2,ws.max_row+1):
        dt=ws.cell(row,1).value
        if not dt: continue
        gross=ws.cell(row,ci['Gross']).value
        key=(dt,rt,str(gross)) if (dt,rt,str(gross)) in g else (dt,rt)
        rid=f"{dt:%Y-%m-%d}_{slug(ws.cell(row,2).value)}_{rt}"+(f"_{gross}" if len(key)==3 else '')
        ws.cell(row,ci['RoundId'],rid)
        RS[key]=(ws,row,ci,rid)
report=[];HOLES=[]
for key,hs in g.items():
    rt=key[1]; m=RS.get(key)
    if m: ws,row,ci,rid=m; course=ws.cell(row,2).value; tees=ws.cell(row,3).value
    else:
        course=hs[0]['Course'];tees=hs[0]['Tees'];rid=f"{key[0]:%Y-%m-%d}_{slug(course)}_{rt}"; report.append(f'NO ROUND ROW: {rid}')
    for h in hs:
        cap=lambda v:{'hit':'Hit','miss':'Miss','NA':'NA','na':'NA'}.get(v,v) if v else None
        HOLES.append([rid,key[0],course,tees,int(rt),int(h['Hole']),int(h['Par']),int(h['Score']) if h['Score'] else None,
          cap(h['FIR']),h['FairwayMiss'] or None,cap(h['GIR']),h['GreenMiss'] or None,int(h['Putts']) if h['Putts'] else None,None,None])
    if not m: continue
    sc=[h for h in hs if h['Score']]; full=len(sc)==len(hs)==int(rt)
    put=lambda c,v,name:(report.append(f'{rid} {name}: {ws.cell(row,c).value} -> {v}') if ws.cell(row,c).value not in (None,'') and ws.cell(row,c).value!=v else None, ws.cell(row,c,v))
    if full:
        tp=[int(h['Score'])-int(h['Par']) for h in hs]; d=sum(x==2 for x in tp); t=sum(x>=3 for x in tp)
        if sum(int(h['Score']) for h in hs)!=ws.cell(row,ci['Gross']).value: report.append(f'{rid} gross mismatch')
        old=(ws.cell(row,ci['Doubles (exact +2)']).value,ws.cell(row,ci['TriplesOrWorse (+3 or more)']).value)
        if old==(None,None): ws.cell(row,ci['Doubles (exact +2)'],d); ws.cell(row,ci['TriplesOrWorse (+3 or more)'],t); report.append(f'{rid} D/T filled {d}/{t}')
        elif old!=(d,t): report.append(f'{rid} D/T existing {old} vs holes {(d,t)} (kept existing)')
    if all(h['GIR'] for h in hs):
        put(ci['GIR'],sum(h['GIR'].lower()=='hit' for h in hs),'GIR')
    if all(h['Putts'] for h in hs):
        p=[int(h['Putts']) for h in hs]; put(ci['Putts'],sum(p),'Putts'); put(ci['ThreePuttsPlus'],sum(x>=3 for x in p),'3P'); put(ci['OnePutts'],sum(x==1 for x in p),'1P')
    f45=[h for h in hs if int(h['Par'])>=4]
    if f45 and all(h['FIR'] and h['FIR']!='NA' for h in f45):
        put(ci['FairwaysHit'],sum(h['FIR']=='hit' for h in f45),'FH'); put(ci['FairwaysPossible'],len(f45),'FP')
# keep Soldier Hollow (and any other) old Fairway Holes rows for dates not in CSV
csvd={k[0] for k in g}
for r in oldFH:
    if r[0] not in csvd:
        m=RS.get((r[0],'18' if r[2]=='18-hole' else '9'))
        HOLES.append([m[3] if m else None,r[0],r[1],m[0].cell(m[1],3).value if m else None,18 if r[2]=='18-hole' else 9,r[3],r[4],r[5],r[7],None,None,None,None,None,None])
        report.append(f'kept old FIR row {r[0]:%Y-%m-%d} {r[1]} H{r[3]}') if r[3]==1 else None
HOLES.sort(key=lambda x:(x[1],x[0] or '',x[5]))
del wb['Holes']; ws=wb.create_sheet('Holes')
ws.append(['RoundId','RoundDate','Course','Tees','RoundType','Hole','Par','Score','ToPar','FIR','FairwayMiss','GIR','GreenMiss','Putts','UpDownAttempt','UpDownMade'])
for c in ws[1]: c.font=Font(bold=True)
for i,h in enumerate(HOLES,2):
    ws.append(h[:8]+[f'=IF(COUNT(G{i}:H{i})=2,H{i}-G{i},"")']+h[8:]); ws.cell(i,2).number_format='yyyy-mm-dd'
ws.freeze_panes='A2'
# regenerate Fairway Holes from Holes (par 4/5 with Hit/Miss), same layout so summary formulas keep working
fh=wb['Fairway Holes']; fh.delete_rows(2,fh.max_row)
i=2
for h in sorted(HOLES,key=lambda x:(x[1],x[5]),reverse=False):
    if h[6]>=4 and h[8] in ('Hit','Miss') and h[7] is not None:
        fh.cell(i,1,h[1]).number_format='yyyy-mm-dd'; fh.cell(i,2,h[2]); fh.cell(i,3,f'{h[4]}-hole'); fh.cell(i,4,h[5]); fh.cell(i,5,h[6]); fh.cell(i,6,h[7])
        fh.cell(i,7,f'=IF(COUNT(E{i}:F{i})=2,F{i}-E{i},"")'); fh.cell(i,8,h[8])
        fh.cell(i,9,f'=IF(C{i}="9-hole",1,IF(COUNTIFS(\'18-Hole Rounds\'!$A$2:$A$1000,A{i},\'18-Hole Rounds\'!$R$2:$R$1000,1)>0,1,0))'); i+=1
wb.save('up/Hunter_Handicap_Tracker.xlsx')
print(len(HOLES),'hole rows;',i-2,'fairway rows'); print('\n'.join(report))
